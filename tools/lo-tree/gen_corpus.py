"""Generate a small pptx corpus that exercises text, fonts, shapes, tables, images, charts, notes."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.dml.color import RGBColor
from PIL import Image, ImageDraw

def text_deck(path):
    prs = Presentation()
    for i, font in enumerate(["Calibri", "Arial", "Times New Roman", "Cambria", "Courier New", "Nonexistent Brand Font"]):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = f"Slide {i+1}: {font}"
        tf = s.placeholders[1].text_frame
        tf.text = "The quick brown fox jumps over the lazy dog. " * 4
        for run in tf.paragraphs[0].runs:
            run.font.name = font; run.font.size = Pt(20)
        p = tf.add_paragraph(); p.text = "Ελληνικά — Кириллица — Ünïcödé — 日本語テキスト"; p.level = 1
        for run in p.runs: run.font.name = font
        s.notes_slide.notes_text_frame.text = f"Speaker notes for slide {i+1}"
    prs.save(path)

def shapes_deck(path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    from pptx.enum.shapes import MSO_SHAPE
    for j, shp in enumerate([MSO_SHAPE.ROUNDED_RECTANGLE, MSO_SHAPE.OVAL, MSO_SHAPE.RIGHT_ARROW, MSO_SHAPE.STAR_5_POINT]):
        sh = s.shapes.add_shape(shp, Inches(0.5 + j*2.2), Inches(1), Inches(2), Inches(1.5))
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(40*j, 90, 200 - 40*j)
        sh.text_frame.text = f"Shape {j}"
    rows, cols = 4, 3
    tbl = s.shapes.add_table(rows, cols, Inches(0.5), Inches(3), Inches(6), Inches(2)).table
    for r in range(rows):
        for c in range(cols): tbl.cell(r, c).text = f"r{r}c{c}"
    img = Image.new("RGB", (400, 300), (250, 240, 200)); d = ImageDraw.Draw(img)
    d.ellipse((50, 50, 350, 250), fill=(20, 120, 220)); img.save("/tmp/pic.png")
    s.shapes.add_picture("/tmp/pic.png", Inches(7), Inches(3), width=Inches(2.5))
    prs.save(path)

def chart_deck(path):
    prs = Presentation()
    for ct in [XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.PIE, XL_CHART_TYPE.LINE_MARKERS]:
        s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = str(ct)
        cd = CategoryChartData(); cd.categories = ["Q1", "Q2", "Q3", "Q4"]
        cd.add_series("2025", (19.2, 21.4, 16.7, 25.0)); cd.add_series("2026", (22.3, 28.6, 15.2, 30.1))
        s.shapes.add_chart(ct, Inches(1), Inches(1.5), Inches(8), Inches(5), cd)
    prs.save(path)

def hidden_deck(path):
    """Three slides; the second is marked hidden (skipped by PDF export and by PNG rendering unless asked)."""
    prs = Presentation()
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = f"Slide {i+1}" + (" (hidden)" if i == 1 else "")
        if i == 1: s._element.set("show", "0")
    prs.save(path)

text_deck("corpus/text_fonts.pptx"); shapes_deck("corpus/shapes_table_image.pptx"); chart_deck("corpus/charts.pptx")
hidden_deck("corpus/hidden.pptx")
print("corpus written")
